from django.shortcuts import render, redirect
from django.contrib import messages
from django.contrib.admin.views.decorators import staff_member_required
from django.core.management import call_command
from django.core.paginator import Paginator
from django.db.models import Count, Q
from django.conf import settings
from django.contrib.auth.models import User
from django.views.decorators.csrf import csrf_protect
from django.views.decorators.http import require_POST
from django.views.decorators.http import require_GET
from django.http import JsonResponse, HttpResponse
from .models import Employee, Department, Position
import os
import tempfile
import json
import csv
import openpyxl
from datetime import datetime

def dashboard(request):
    from django.utils import timezone
    from datetime import timedelta
    from django.db.models import Count, Q, Case, When, IntegerField, Sum
    
    # Base queryset - hanya karyawan aktif (tgl_out kosong)
    active_employees = Employee.objects.filter(tgl_out__isnull=True)
    
    # Statistik - sama persis dengan di employee_list
    stats = {
        'total': active_employees.count(),
        'tetap': active_employees.filter(status_karyawan='tetap').count(),
        'kontrak': active_employees.filter(status_karyawan='kontrak').count(),
        'os': active_employees.filter(status_karyawan='os').count(),
    }
    
    # Karyawan baru bulan ini (hanya 3)
    first_day_of_month = timezone.now().replace(day=1, hour=0, minute=0, second=0, microsecond=0)
    new_employees_this_month = active_employees.filter(
        tgl_rekrut__gte=first_day_of_month
    ).order_by('-tgl_rekrut')[:3]
    
    # SEBARAN KARYAWAN PER DEPARTMENT + STATUS
    employees_by_dept = []
    
    # Ambil semua department yang punya karyawan
    departments = Department.objects.filter(employees__tgl_out__isnull=True).distinct()
    
    for dept in departments:
        # Hitung karyawan per status di department ini
        tetap_count = active_employees.filter(department=dept, status_karyawan='tetap').count()
        kontrak_count = active_employees.filter(department=dept, status_karyawan='kontrak').count()
        os_count = active_employees.filter(department=dept, status_karyawan='os').count()
        total = tetap_count + kontrak_count + os_count
        
        if total > 0:  # Hanya tampilkan yang punya karyawan
            employees_by_dept.append({
                'name': dept.name,
                'total': total,
                'tetap': tetap_count,
                'kontrak': kontrak_count,
                'os': os_count,
            })
    
    # Sort by total DESC
    employees_by_dept.sort(key=lambda x: x['total'], reverse=True)
    
    # Data untuk chart status (tetap/kontrak/os)
    employees_by_status = active_employees.values('status_karyawan').annotate(count=Count('id'))
    
    context = {
        'stats': stats,
        'employees_by_dept': employees_by_dept,
        'employees_by_status': employees_by_status,
        'new_employees': new_employees_this_month,
        'total_departments': Department.objects.count(),
        'total_positions': Position.objects.count(),
    }
    return render(request, 'hr/dashboard.html', context)

@staff_member_required
def import_employees_view(request):
    if request.method == 'POST' and request.FILES.get('file'):
        file = request.FILES['file']
        
        # Simpan file sementara
        with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as tmp:
            for chunk in file.chunks():
                tmp.write(chunk)
            tmp_path = tmp.name
        
        try:
            # Panggil management command
            from io import StringIO
            output = StringIO()
            call_command('import_employees', tmp_path, stdout=output)
            
            # Tampilkan output ke user
            messages.success(request, f'Import berhasil!\n\n{output.getvalue()}')
        except Exception as e:
            messages.error(request, f'Error: {str(e)}')
        finally:
            # Hapus file sementara
            os.unlink(tmp_path)
        
        return redirect('admin:hr_employee_changelist')
    
    return redirect('admin:hr_employee_changelist')

@staff_member_required
def employee_list(request):
    """
    View utama untuk menampilkan daftar karyawan dengan semua fitur:
    - Hanya menampilkan karyawan aktif (tgl_out kosong)
    - Pencarian real-time
    - Filter status
    - Sorting
    - Pagination
    - Card/Table view
    """
    # Ambil parameter dari request
    search_query = request.GET.get('search', '').strip()
    status_filter = request.GET.get('status_karyawan', '')
    department_filter = request.GET.get('department', '')
    position_filter = request.GET.get('position', '')
    sort_by = request.GET.get('sort', 'employee_id')
    order = request.GET.get('order', 'asc')
    view_mode = request.GET.get('view', 'table')  # 'table' atau 'card'
    
    # BASE QUERYSET - HANYA KARYAWAN AKTIF (tgl_out kosong)
    employees = Employee.objects.filter(tgl_out__isnull=True).select_related('department', 'position')
    
    # Filter berdasarkan pencarian
    if search_query:
        employees = employees.filter(
            Q(employee_id__icontains=search_query) |
            Q(nama__icontains=search_query) |
            Q(no_ktp__icontains=search_query) |
            Q(no_hp__icontains=search_query)
        )
    
    # Filter berdasarkan status
    if status_filter:
        # Filter berdasarkan status_karyawan (Tetap/Kontrak/OS)
        employees = employees.filter(status_karyawan=status_filter)
    
    # Filter berdasarkan department
    if department_filter:
        employees = employees.filter(department_id=department_filter)
    
    # Filter berdasarkan position
    if position_filter:
        employees = employees.filter(position_id=position_filter)
    
    # Sorting
    if order == 'desc':
        sort_by = f'-{sort_by}'
    employees = employees.order_by(sort_by)
    
    # STATISTIK - HANYA KARYAWAN AKTIF
    stats = {
        'total': Employee.objects.filter(tgl_out__isnull=True).count(),
        'tetap': Employee.objects.filter(tgl_out__isnull=True, status_karyawan='tetap').count(),
        'kontrak': Employee.objects.filter(tgl_out__isnull=True, status_karyawan='kontrak').count(),
        'os': Employee.objects.filter(tgl_out__isnull=True, status_karyawan='os').count(),
    }
    
    # Data untuk filter dropdown
    departments = Department.objects.filter(is_active=True).values('id', 'name')
    positions = Position.objects.filter(is_active=True).values('id', 'title')
    
    # Pagination
    paginator = Paginator(employees, 9)
    page_number = request.GET.get('page')
    page_obj = paginator.get_page(page_number)
    
    # Siapkan parameter untuk maintain filter
    params = {
        'search': search_query,
        'status': status_filter,
        'department': department_filter,
        'position': position_filter,
        'sort': sort_by.lstrip('-'),
        'order': order,
        'view': view_mode,
    }
    
    context = {
        'page_obj': page_obj,
        'stats': stats,
        'departments': departments,
        'positions': positions,
        'params': params,
        'view_mode': view_mode,
        'search_query': search_query,
        'status_filter': status_filter,
        'department_filter': int(department_filter) if department_filter else None,
        'position_filter': int(position_filter) if position_filter else None,
        'sort_by': sort_by.lstrip('-'),
        'order': order,
    }
    
    return render(request, 'hr/employee_list.html', context)

@staff_member_required
@require_POST
def employee_search_api(request):
    """
    API endpoint untuk pencarian real-time dengan debounce
    Mengembalikan JSON untuk update tabel tanpa reload
    """
    data = json.loads(request.body)
    search_query = data.get('search', '').strip()
    status_filter = data.get('status', '')
    department_filter = data.get('department', '')
    position_filter = data.get('position', '')
    page = int(data.get('page', 1))
    
    # BASE QUERYSET - HANYA KARYAWAN AKTIF
    employees = Employee.objects.filter(tgl_out__isnull=True).select_related('department', 'position')
    
    if search_query:
        employees = employees.filter(
            Q(employee_id__icontains=search_query) |
            Q(nama__icontains=search_query)
        )
    
    if status_filter:
        employees = employees.filter(employment_status=status_filter)
    
    if department_filter:
        employees = employees.filter(department_id=department_filter)
    
    if position_filter:
        employees = employees.filter(position_id=position_filter)
    
    paginator = Paginator(employees, 25)
    page_obj = paginator.get_page(page)
    
    # Format data untuk JSON
    employees_data = []
    for emp in page_obj:
        employees_data.append({
            'id': emp.id,
            'employee_id': emp.employee_id,
            'nama': emp.nama,
            'department': emp.department.name if emp.department else '-',
            'position': emp.position.title if emp.position else '-',
            'employment_status': emp.employment_status,
            'status_karyawan': emp.status_karyawan,
            'avatar_color': emp.department.name if emp.department else '',
            'initials': get_initials(emp.nama),
        })
    
    return JsonResponse({
        'employees': employees_data,
        'total_pages': paginator.num_pages,
        'current_page': page_obj.number,
        'total_count': paginator.count,
    })

def get_initials(nama):
    """Helper function untuk mendapatkan inisial dari nama"""
    if not nama:
        return '??'
    parts = nama.split()
    if len(parts) >= 2:
        return (parts[0][0] + parts[1][0]).upper()
    return nama[:2].upper()

@staff_member_required
def export_employees(request):
    """
    Export data karyawan ke Excel atau CSV
    Mendukung format pilihan dan kolom yang dipilih
    """
    export_format = request.GET.get('format', 'excel')
    selected_columns = request.GET.getlist('columns')
    status_filter = request.GET.get('status', '')
    department_filter = request.GET.get('department', '')
    
    # Default columns jika tidak dipilih
    if not selected_columns:
        selected_columns = ['employee_id', 'nama', 'department', 'position', 
                           'employment_status', 'no_hp', 'email']
    
    # BASE QUERYSET - HANYA KARYAWAN AKTIF
    employees = Employee.objects.filter(tgl_out__isnull=True).select_related('department', 'position')
    
    if status_filter:
        employees = employees.filter(employment_status=status_filter)
    if department_filter:
        employees = employees.filter(department_id=department_filter)
    
    # Mapping field ke display name
    field_display = {
        'employee_id': 'NIK',
        'nama': 'Nama Lengkap',
        'gender': 'Jenis Kelamin',
        'tgl_lahir': 'Tanggal Lahir',
        'tempat_lahir': 'Tempat Lahir',
        'no_ktp': 'No KTP',
        'no_kk': 'No KK',
        'no_hp': 'No HP',
        'alamat': 'Alamat',
        'kelurahan': 'Kelurahan',
        'kecamatan': 'Kecamatan',
        'kabupaten_kota': 'Kabupaten/Kota',
        'provinsi': 'Provinsi',
        'kode_pos': 'Kode Pos',
        'status_kawin': 'Status Kawin',
        'agama': 'Agama',
        'pendidikan': 'Pendidikan',
        'department': 'Department',
        'position': 'Posisi',
        'employment_status': 'Status',
        'status_karyawan': 'Status Karyawan',
        'tgl_rekrut': 'Tanggal Rekrut',
        'tgl_out': 'Tanggal Keluar',
        'no_rek_bank': 'No Rekening',
        'nama_bank': 'Nama Bank',
        'bpjs_tk_no': 'No BPJS TK',
        'bpjs_kes_no': 'No BPJS Kesehatan',
        'no_npwp': 'NPWP',
    }
    
    if export_format == 'excel':
        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename="employees_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx"'
        
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Employees"
        
        # Header
        for col, field in enumerate(selected_columns, 1):
            ws.cell(row=1, column=col, value=field_display.get(field, field))
        
        # Data
        for row, emp in enumerate(employees, 2):
            for col, field in enumerate(selected_columns, 1):
                value = getattr(emp, field, '')
                if field == 'department' and emp.department:
                    value = emp.department.name
                elif field == 'position' and emp.position:
                    value = emp.position.title
                elif field == 'employment_status':
                    value = dict(Employee.EMPLOYMENT_STATUS).get(emp.employment_status, '')
                elif field == 'status_karyawan':
                    value = dict(Employee.STATUS_KARYAWAN).get(emp.status_karyawan, '')
                ws.cell(row=row, column=col, value=value)
        
        wb.save(response)
        return response
    
    else:  # CSV
        response = HttpResponse(content_type='text/csv')
        response['Content-Disposition'] = f'attachment; filename="employees_{datetime.now().strftime("%Y%m%d_%H%M%S")}.csv"'
        
        writer = csv.writer(response)
        # Header
        writer.writerow([field_display.get(field, field) for field in selected_columns])
        
        # Data
        for emp in employees:
            row = []
            for field in selected_columns:
                value = getattr(emp, field, '')
                if field == 'department' and emp.department:
                    value = emp.department.name
                elif field == 'position' and emp.position:
                    value = emp.position.title
                elif field == 'employment_status':
                    value = dict(Employee.EMPLOYMENT_STATUS).get(emp.employment_status, '')
                elif field == 'status_karyawan':
                    value = dict(Employee.STATUS_KARYAWAN).get(emp.status_karyawan, '')
                row.append(value)
            writer.writerow(row)
        
        return response

@staff_member_required
def get_filter_options(request):
    """API untuk mendapatkan opsi filter (department, position)"""
    departments = list(Department.objects.filter(is_active=True).values('id', 'name'))
    positions = list(Position.objects.filter(is_active=True).values('id', 'title'))
    
    return JsonResponse({
        'departments': departments,
        'positions': positions,
    })

@staff_member_required
@csrf_protect
def create_user_view(request):
    if request.method == 'POST':
        username = request.POST.get('username')
        password = request.POST.get('password')
        first_name = request.POST.get('first_name', '')
        last_name = request.POST.get('last_name', '')
        email = request.POST.get('email', '')
        employee_id = request.POST.get('employee_id', '')
        is_superuser = request.POST.get('is_superuser') == 'on'
        
        if User.objects.filter(username=username).exists():
            messages.error(request, f'Username {username} sudah ada')
        else:
            if is_superuser:
                user = User.objects.create_superuser(
                    username=username,
                    email=email,
                    password=password,
                    first_name=first_name,
                    last_name=last_name
                )
            else:
                user = User.objects.create_user(
                    username=username,
                    email=email,
                    password=password,
                    first_name=first_name,
                    last_name=last_name
                )
            
            if employee_id:
                try:
                    employee = Employee.objects.get(employee_id=employee_id)
                    employee.user = user
                    employee.save()
                    messages.success(request, f'User {username} dibuat dan terhubung ke employee {employee_id}')
                except Employee.DoesNotExist:
                    messages.warning(request, f'User dibuat tapi employee {employee_id} tidak ditemukan')
            else:
                messages.success(request, f'User {username} berhasil dibuat')
        
        return redirect('admin:auth_user_changelist')
    
    return redirect('admin:auth_user_changelist')


@require_GET
def employee_detail_api(request, employee_id):
    """API untuk mengambil detail karyawan"""
    try:
        employee = Employee.objects.select_related(
            'department', 'position', 'user'
        ).get(employee_id=employee_id)
        
        # Siapkan data untuk dikembalikan
        data = {
            'id': employee.id,
            'employee_id': employee.employee_id,
            'nama': employee.nama,
            'gender': employee.get_gender_display() if employee.gender else '-',
            'tgl_lahir': employee.tgl_lahir.strftime('%d %B %Y') if employee.tgl_lahir else '-',
            'tempat_lahir': employee.tempat_lahir or '-',
            'no_ktp': employee.no_ktp or '-',
            'no_kk': employee.no_kk or '-',
            'no_hp': employee.no_hp or '-',
            'alamat': employee.alamat or '-',
            'kelurahan': employee.kelurahan or '-',
            'kecamatan': employee.kecamatan or '-',
            'kabupaten_kota': employee.kabupaten_kota or '-',
            'provinsi': employee.provinsi or '-',
            'kode_pos': employee.kode_pos or '-',
            'status_kawin': employee.get_status_kawin_display() if employee.status_kawin else '-',
            'tanggungan': employee.tanggungan or 0,
            'agama': employee.get_agama_display() if employee.agama else '-',
            'pendidikan': employee.pendidikan or '-',
            
            # Data Kepegawaian
            'tgl_rekrut': employee.tgl_rekrut.strftime('%d %B %Y') if employee.tgl_rekrut else '-',
            'status_karyawan': employee.get_status_karyawan_display() if employee.status_karyawan else '-',
            'tgl_kartetap': employee.tgl_kartetap.strftime('%d %B %Y') if employee.tgl_kartetap else '-',
            'posisi_karyawan': employee.posisi_karyawan or '-',
            'department': employee.department.name if employee.department else '-',
            'position': employee.position.title if employee.position else '-',
            'group': employee.group or '-',
            'dept': employee.dept or '-',
            'jabatan': employee.jabatan or '-',
            'kontrak_ke': employee.kontrak_ke or 0,
            'kontrak_berakhir': employee.kontrak_berakhir.strftime('%d %B %Y') if employee.kontrak_berakhir else '-',
            'kode_gaji': employee.kode_gaji or '-',
            
            # Bank & BPJS
            'no_rek_bank': employee.no_rek_bank or '-',
            'nama_bank': employee.nama_bank or '-',
            'bpjs_tk_no': employee.bpjs_tk_no or '-',
            'bpjs_kes_no': employee.bpjs_kes_no or '-',
            'no_npwp': employee.no_npwp or '-',
            
            # Status
            'employment_status': employee.get_employment_status_display() if employee.employment_status else '-',
            'tgl_out': employee.tgl_out.strftime('%d %B %Y') if employee.tgl_out else '-',
        }
        
        return JsonResponse({
            'success': True,
            'data': data
        })
    except Employee.DoesNotExist:
        return JsonResponse({
            'success': False,
            'error': 'Karyawan tidak ditemukan'
        }, status=404)